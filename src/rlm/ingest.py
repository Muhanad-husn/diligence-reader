"""Reads a sample's data room into runs/<sample>/sections.jsonl.

One reader per format. PDF text comes from pdftext, page by page, and every page is extracted
a second time with pypdf as an independent check; a page the two engines disagree on is written
with a warning and counted, never repaired. Ruled tables come from pdfplumber, one section per
table row. Workbooks come from openpyxl with data_only=True, one section per non-empty row and
no header inference. CSV, plain text, markdown, EML and MBOX come from the standard library.

Records are sorted by doc then ordinal and their JSON keys are sorted, so two runs of the same
sample write the same bytes. Nothing here reads a key, opens a socket or calls a model.
"""

from __future__ import annotations

import csv
import datetime
import email
import email.policy
import json
import mailbox
import re
import sys
from collections import Counter
from collections.abc import Callable, Iterable
from dataclasses import dataclass, replace
from pathlib import Path

import openpyxl
import pdfplumber
import pypdf
from openpyxl.utils import get_column_letter
from pdftext.extraction import paginated_plain_text_output

from rlm.sections import (
    Cell,
    Section,
    line_anchor,
    message_line_anchor,
    page_line_anchor,
    page_table_row_anchor,
    row_anchor,
    sheet_cell_anchor,
    to_record,
)
from rlm.index import build_index, write_index

EXTENSIONS = frozenset({".pdf", ".xlsx", ".csv", ".txt", ".md", ".eml", ".mbox"})

# The mail headers that become a mail section's heading, in this order.
MAIL_HEADERS = ("From", "To", "Date", "Subject")

# A block of one line that starts with a section number, such as "3. Scope of exposure", is the
# heading of the blocks that follow it.
_HEADING = re.compile(r"^\d+(?:\.\d+)*\.?\s+\S")

# pypdf renders the bullet these documents use at U+007F where pdftext renders it at U+2022.
# The cross-check maps the one to the other before it counts words, so it compares words rather
# than that one glyph.
_PYPDF_BULLET = "\x7f"
_BULLET = "•"

_BLANK = re.compile(r"^\s*$")


@dataclass(frozen=True)
class Coverage:
    """What one run read: documents, sections written, disagreeing pages and empty parts.

    An empty part is a PDF page one of the two engines returned no text for, or a document of
    any format that yielded no section at all.
    """

    documents: int
    sections: int
    disagreements: int
    empty: int


def coverage_line(coverage: Coverage) -> str:
    """Writes the one line a run prints when it finishes."""
    return (
        f"documents read: {coverage.documents}, "
        f"sections written: {coverage.sections}, "
        f"engine disagreements: {coverage.disagreements}, "
        f"empty extractions: {coverage.empty}"
    )


def split_blocks(lines: list[str]) -> list[tuple[int, str]]:
    """Splits lines on blank lines and returns each block with its first line number, 1-based."""
    blocks = []
    start = None
    for number, line in enumerate(lines, start=1):
        if _BLANK.match(line):
            if start is not None:
                blocks.append((start, "\n".join(lines[start - 1 : number - 1])))
                start = None
        elif start is None:
            start = number
    if start is not None:
        blocks.append((start, "\n".join(lines[start - 1 :])))
    return blocks


def serialise(value) -> str | int | float:
    """Turns a workbook value into the form a record holds.

    A date, time or timestamp becomes its ISO string, a number stays a number, and everything
    else becomes text.
    """
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    return str(value)


def text_sections(
    doc: str,
    lines: list[str],
    anchor_for: Callable[[int], str],
    heading: str | None,
    warning: str | None,
) -> tuple[list[Section], str | None]:
    """Makes one section per block of lines, and returns them with the heading left in force."""
    sections = []
    for first, block in split_blocks(lines):
        if "\n" not in block and _HEADING.match(block):
            heading = block
        sections.append(
            Section(
                doc=doc,
                ordinal=0,
                kind="text",
                anchor=anchor_for(first),
                heading=heading,
                text=block,
                warning=warning,
            )
        )
    return sections, heading


def read_pdf(path: Path, doc: str) -> tuple[list[Section], int, int]:
    """Reads a PDF into text sections and table rows, and counts its pages the engines differ on.

    pdftext gives the text of each page, pypdf gives it again, and a page whose word multisets
    differ or whose extraction is empty carries a warning naming the difference. The pdftext
    text is the record either way.
    """
    pages = paginated_plain_text_output(str(path))
    others = [page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages]
    with pdfplumber.open(path) as plumbed:
        tables_by_page = [page.extract_tables() for page in plumbed.pages]

    sections: list[Section] = []
    heading: str | None = None
    disagreements = 0
    empty = 0
    for number, text in enumerate(pages, start=1):
        other = others[number - 1].replace(_PYPDF_BULLET, _BULLET)
        notes = []
        if not text.strip():
            notes.append("pdftext extracted no text")
        if not other.strip():
            notes.append("pypdf extracted no text")
        if notes:
            empty += 1
        mine, theirs = Counter(text.split()), Counter(other.split())
        if mine != theirs:
            disagreements += 1
            notes.append(
                f"{sum((mine - theirs).values())} words only in pdftext and "
                f"{sum((theirs - mine).values())} only in pypdf"
            )
        warning = f"page {number}: " + "; ".join(notes) if notes else None

        page_sections, heading = text_sections(
            doc,
            text.split("\n"),
            lambda line, page=number: page_line_anchor(doc, page, line),
            heading,
            warning,
        )
        sections.extend(page_sections)
        for table_number, table in enumerate(tables_by_page[number - 1], start=1):
            for row_number, row in enumerate(table, start=1):
                cells = tuple(
                    Cell(ref=f"t{table_number}r{row_number}c{column}", value=value or "")
                    for column, value in enumerate(row, start=1)
                )
                if not any(str(cell.value).strip() for cell in cells):
                    continue
                sections.append(
                    Section(
                        doc=doc,
                        ordinal=0,
                        kind="row",
                        anchor=page_table_row_anchor(doc, number, table_number, row_number),
                        heading=heading,
                        text=" | ".join(str(cell.value) for cell in cells),
                        warning=warning,
                        cells=cells,
                    )
                )
    return sections, disagreements, empty


def read_workbook(path: Path, doc: str) -> list[Section]:
    """Reads a workbook into one row section per non-empty row of every sheet.

    Cell values are the cached ones, so a formula reads as the number the file holds. No row is
    treated as a header.
    """
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sections = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            entries = [
                (get_column_letter(cell.column), cell.row, serialise(cell.value))
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            ]
            if not entries:
                continue
            column, number, _ = entries[0]
            sections.append(
                Section(
                    doc=doc,
                    ordinal=0,
                    kind="row",
                    anchor=sheet_cell_anchor(doc, sheet.title, column, number),
                    heading=sheet.title,
                    text=" | ".join(str(value) for _, _, value in entries),
                    warning=None,
                    cells=tuple(Cell(ref=f"{letter}{at}", value=value) for letter, at, value in entries),
                )
            )
    workbook.close()
    return sections


def read_csv(path: Path, doc: str) -> list[Section]:
    """Reads a CSV file into one row section per row, with no sniffing and no header rule."""
    sections = []
    with path.open(newline="", encoding="utf-8") as handle:
        for number, row in enumerate(csv.reader(handle), start=1):
            cells = tuple(
                Cell(ref=f"r{number}c{column}", value=value)
                for column, value in enumerate(row, start=1)
            )
            if not any(cell.value.strip() for cell in cells):
                continue
            sections.append(
                Section(
                    doc=doc,
                    ordinal=0,
                    kind="row",
                    anchor=row_anchor(doc, number),
                    heading=None,
                    text=" | ".join(row),
                    warning=None,
                    cells=cells,
                )
            )
    return sections


def read_text(path: Path, doc: str) -> list[Section]:
    """Reads a plain text or markdown file into one section per block of lines."""
    lines = path.read_text(encoding="utf-8").split("\n")
    sections, _ = text_sections(doc, lines, lambda line: line_anchor(doc, line), None, None)
    return sections


def mail_sections(doc: str, message, anchor_for: Callable[[int], str]) -> list[Section]:
    """Makes one section per block of a mail body, all headed by the message's own headers."""
    heading = " | ".join(f"{name}: {message[name]}" for name in MAIL_HEADERS if message[name])
    body = message.get_body(preferencelist=("plain",))
    text = body.get_content() if body is not None else ""
    return [
        Section(
            doc=doc,
            ordinal=0,
            kind="text",
            anchor=anchor_for(first),
            heading=heading or None,
            text=block,
            warning=None,
        )
        for first, block in split_blocks(text.split("\n"))
    ]


def read_eml(path: Path, doc: str) -> list[Section]:
    """Reads one mail message into sections anchored by line number within its body."""
    message = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    return mail_sections(doc, message, lambda line: line_anchor(doc, line))


def read_mbox(path: Path, doc: str) -> list[Section]:
    """Reads an mbox into sections anchored by message ordinal and line number within its body."""
    box = mailbox.mbox(
        str(path),
        factory=lambda handle: email.message_from_binary_file(handle, policy=email.policy.default),
    )
    sections = []
    for number, message in enumerate(box, start=1):
        sections.extend(
            mail_sections(doc, message, lambda line, at=number: message_line_anchor(doc, at, line))
        )
    box.close()
    return sections


def read_document(path: Path, doc: str) -> tuple[list[Section], int, int]:
    """Reads one file with the reader for its format and numbers its sections from 1."""
    suffix = path.suffix.lower()
    disagreements = 0
    empty = 0
    if suffix == ".pdf":
        sections, disagreements, empty = read_pdf(path, doc)
    elif suffix == ".xlsx":
        sections = read_workbook(path, doc)
    elif suffix == ".csv":
        sections = read_csv(path, doc)
    elif suffix in (".txt", ".md"):
        sections = read_text(path, doc)
    elif suffix == ".eml":
        sections = read_eml(path, doc)
    elif suffix == ".mbox":
        sections = read_mbox(path, doc)
    else:
        raise ValueError(f"no reader for {path}")
    if not sections:
        empty += 1
    return (
        [replace(section, ordinal=ordinal) for ordinal, section in enumerate(sections, start=1)],
        disagreements,
        empty,
    )


def documents(sample_dir: Path) -> list[Path]:
    """Lists the files of a sample's data room this phase reads, in path order."""
    room = sample_dir / "data_room"
    found = [path for path in room.rglob("*") if path.is_file() and path.suffix.lower() in EXTENSIONS]
    return sorted(found, key=lambda path: path.relative_to(sample_dir).as_posix())


def write_records(run_dir: Path, records: Iterable[dict]) -> None:
    """Writes one JSON object per line, keys sorted, lines ending in a single newline."""
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "sections.jsonl").write_text(
        "".join(json.dumps(record, sort_keys=True, ensure_ascii=False) + "\n" for record in records),
        encoding="utf-8",
        newline="\n",
    )


def ingest(sample_dir: Path, run_dir: Path) -> Coverage:
    """Reads a sample's data room, writes run_dir / "sections.jsonl" and returns the coverage."""
    records = []
    read = 0
    disagreements = 0
    empty = 0
    for path in documents(sample_dir):
        sections, disagreed, blank = read_document(path, path.relative_to(sample_dir).as_posix())
        read += 1
        disagreements += disagreed
        empty += blank
        records.extend(to_record(section) for section in sections)
    records.sort(key=lambda record: (record["doc"], record["ordinal"]))
    write_records(run_dir, records)
    write_index(run_dir, build_index(records))
    return Coverage(
        documents=read, sections=len(records), disagreements=disagreements, empty=empty
    )


def main(argv: list[str]) -> int:
    """Ingests one sample from the command line and prints the coverage line."""
    if len(argv) != 2:
        print("usage: python -m rlm.ingest <sample_dir> <run_dir>")
        return 2
    sample_dir, run_dir = argv
    print(coverage_line(ingest(Path(sample_dir), Path(run_dir))))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
