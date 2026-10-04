"""Phase 8 slice 01: report.md exported to docx, pdf and evidence.csv.

The pinned report of each gate sample is copied into a temporary directory and exported there,
so nothing is written under runs/. Recall is the key's own measure, read from each format's text.
"""

import csv
import io
import re
import subprocess
import sys
from pathlib import Path

import pytest
from pypdf import PdfReader

from conftest import ROOT
from rlm.grade import measure_recall
from rlm.key import load_key

# Facts the report body carries that the Evidence section does not. price-reduction is the deal
# price cut the writer computes; it is not a quote from the room.
EVIDENCE_ABSENT = {"atlas": {"price-reduction"}, "northwind": set(), "northstar-dental": set()}


def evidence_section(report: str) -> str:
    """The text under the report's `## Evidence` heading, to the next second level heading."""
    match = re.search(r"^## Evidence\s*$", report, re.M)
    assert match, "report has no Evidence section"
    rest = report[match.end():]
    following = re.search(r"^## ", rest, re.M)
    return rest[: following.start()] if following else rest


def bullet_count(report: str) -> int:
    return sum(1 for line in evidence_section(report).splitlines() if line.startswith("- "))


def docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


@pytest.fixture
def report(sample):
    path = ROOT / "runs" / sample / "report.md"
    if not path.exists():
        pytest.skip(f"no pinned report for {sample}")
    return path.read_text(encoding="utf-8")


def exported(tmp_path: Path, name: str, report: str) -> Path:
    out = tmp_path / name
    out.mkdir()
    (out / "report.md").write_text(report, encoding="utf-8", newline="")
    from rlm.export import export

    export(out)
    return out


@pytest.fixture
def first(tmp_path, report):
    return exported(tmp_path, "one", report)


def recall(sample_dir, text):
    return measure_recall(load_key(sample_dir), text)


def test_docx_recall_equals_report(sample_dir, report, first):
    assert recall(sample_dir, docx_text(first / "report.docx"))[0] == recall(sample_dir, report)[0] == 100


def test_pdf_recall_equals_report(sample_dir, report, first):
    assert recall(sample_dir, pdf_text(first / "report.pdf"))[0] == recall(sample_dir, report)[0] == 100


def test_csv_recall_equals_evidence_section(sample, sample_dir, report, first):
    text = (first / "evidence.csv").read_text(encoding="utf-8-sig")
    _, from_csv, _ = recall(sample_dir, text)
    _, from_section, _ = recall(sample_dir, evidence_section(report))
    assert set(from_csv) == set(from_section)
    _, from_report, _ = recall(sample_dir, report)
    missing = set(from_report) - set(from_csv)
    assert missing == EVIDENCE_ABSENT[sample], f"facts in report.md but not in evidence.csv: {sorted(missing)}"


def test_csv_shape(report, first):
    raw = (first / "evidence.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    body = raw[3:]
    assert body.endswith(b"\r\n")
    assert body.count(b"\n") == body.count(b"\r\n"), "a bare line feed outside CRLF"
    rows = list(csv.reader(io.StringIO(body.decode("utf-8"), newline="")))
    assert rows[0] == ["document", "date", "quote", "citation"]
    assert len(rows) == bullet_count(report) + 1


def test_two_exports_agree(tmp_path, report, first):
    second = exported(tmp_path, "two", report)
    assert (first / "evidence.csv").read_bytes() == (second / "evidence.csv").read_bytes()
    assert docx_text(first / "report.docx") == docx_text(second / "report.docx")
    assert pdf_text(first / "report.pdf") == pdf_text(second / "report.pdf")


def test_dollar_amounts_survive(report, first):
    amounts = set(re.findall(r"\$\d[\d,.]*[mMkKbB]?", report))
    assert amounts, "the sample has no dollar amounts to check"
    text = docx_text(first / "report.docx")
    lost = sorted(a for a in amounts if a not in text)
    assert not lost, f"dollar amounts lost from the docx: {lost[:10]}"


def test_pdf_keeps_dollar_amounts(report, first):
    text = re.sub(r"\s+", "", pdf_text(first / "report.pdf"))
    amounts = set(re.findall(r"\$\d[\d,.]*[mMkKbB]?", report))
    lost = sorted(a for a in amounts if a not in text)
    assert not lost, f"dollar amounts lost from the pdf: {lost[:10]}"


def test_bullet_parser_standard_shape():
    from rlm.export import parse_bullet

    line = '- 2025-10-18 | "a | b" | [DR-069 | data_room/x.docx#p1]'
    assert parse_bullet(line) == ("2025-10-18", "a | b", "DR-069 | data_room/x.docx#p1")


def test_bullet_parser_comparison_shape():
    from rlm.export import parse_bullet

    line = '- DR-004 says "x" [DR-004 | a#b] against DR-005 says "y" [DR-005 | c#d]'
    assert parse_bullet(line) == ("", line[2:], "")


def test_command_writes_three_files(tmp_path, report):
    (tmp_path / "report.md").write_text(report, encoding="utf-8", newline="")
    done = subprocess.run(
        [sys.executable, "-m", "rlm.export", str(tmp_path)], capture_output=True, text=True
    )
    assert done.returncode == 0, done.stderr
    for name in ("report.docx", "report.pdf", "evidence.csv"):
        assert (tmp_path / name).stat().st_size > 0
