import os, re, csv, sys
import openpyxl
import fitz  # PyMuPDF

ROOT = r"C:\Users\johna\OneDrive\Documents\DATA ROOM TEST SONNET\data_room"
OUT = r"C:\Users\johna\OneDrive\Documents\DATA ROOM TEST SONNET\corpus.txt"

# ---- 1. Load index to build DR-ID mapping ----
idx_path = os.path.join(ROOT, "00_Index_and_Process", "Data_Room_Index.xlsx")
wb = openpyxl.load_workbook(idx_path, data_only=True)
ws = wb["Index"]
rows = list(ws.iter_rows(values_only=True))
# find header row
hdr_i = None
for i, r in enumerate(rows):
    if r and r[0] == "Doc ID":
        hdr_i = i
        break
index_rows = []  # (id, title, folder)
for r in rows[hdr_i+1:]:
    if r and r[0] and str(r[0]).startswith("DR-"):
        index_rows.append((str(r[0]).strip(), str(r[1]).strip(), str(r[2]).strip()))

FOLDER_MAP = {
    "Index and Process": "00_Index_and_Process",
    "Corporate and Board": "01_Corporate_and_Board",
    "Financials and Tax": "02_Financials_and_Tax",
    "Commercial and Advertising": "03_Commercial_and_Advertising",
    "Product Data and Technology": "04_Product_Data_and_Technology",
    "Security IT and Infrastructure": "05_Security_IT_and_Infrastructure",
    "Legal Regulatory and Compliance": "06_Legal_Regulatory_and_Compliance",
    "HR Operations and Integration": "07_HR_Operations_and_Integration",
}

STOP = set("the and of a to in for on with by amp project vistaport media inc draft summary report schedule plan memo overview review redacted final".split())
def toks(s):
    s = s.lower().replace("&", " ").replace("-", " ").replace("_", " ").replace(".", " ")
    return set(w for w in re.split(r"[^a-z0-9]+", s) if w and w not in STOP and len(w) > 1)

# map each file to best DR row within its folder
files_by_folder = {}
for fol in set(FOLDER_MAP.values()):
    p = os.path.join(ROOT, fol)
    files_by_folder[fol] = sorted(os.listdir(p))

assignments = {}  # filepath -> (id, title)
used = set()
for (did, title, folder) in index_rows:
    folpath = FOLDER_MAP.get(folder)
    if not folpath:
        continue
    cand = files_by_folder.get(folpath, [])
    tt = toks(title)
    best, bestscore = None, -1
    for f in cand:
        if (folpath, f) in used:
            continue
        ft = toks(f)
        score = len(tt & ft)
        if score > bestscore:
            bestscore, best = score, f
    if best is not None:
        used.add((folpath, best))
        assignments[os.path.join(ROOT, folpath, best)] = (did, title)

# report unmatched files
all_files = []
for fol, fl in files_by_folder.items():
    for f in fl:
        if f == "README.md":
            continue
        all_files.append(os.path.join(ROOT, fol, f))
unmatched = [f for f in all_files if f not in assignments]

# ---- 2. Extractors ----
def extract_pdf(path):
    doc = fitz.open(path)
    out = []
    for pg in doc:
        out.append(pg.get_text())
    doc.close()
    return "\n".join(out)

def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- worksheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells))
    return "\n".join(out)

def extract_text(path):
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()

def extract(path):
    ext = path.lower().rsplit(".", 1)[-1]
    try:
        if ext == "pdf":
            return extract_pdf(path)
        if ext == "xlsx":
            return extract_xlsx(path)
        return extract_text(path)  # csv, txt, eml, mbox
    except Exception as e:
        return f"[EXTRACTION ERROR: {e}]"

# ---- 3. Write corpus in DR order ----
# build id -> filepath
id_to_path = {}
for path, (did, title) in assignments.items():
    id_to_path[did] = (path, title)

total_chars = 0
with open(OUT, "w", encoding="utf-8") as out:
    for (did, title, folder) in index_rows:
        if did not in id_to_path:
            out.write(f"\n\n{'='*90}\n{did} | {title} | [FILE NOT MATCHED]\n{'='*90}\n")
            continue
        path, _ = id_to_path[did]
        fname = os.path.basename(path)
        text = extract(path)
        total_chars += len(text)
        out.write(f"\n\n{'='*90}\n{did} | {title}\nfolder: {folder} | file: {fname}\n{'='*90}\n")
        out.write(text)

print("INDEX ROWS:", len(index_rows))
print("MATCHED:", len(assignments))
print("UNMATCHED FILES:", unmatched)
print("TOTAL TEXT CHARS:", total_chars)
print("CORPUS WRITTEN:", OUT)
